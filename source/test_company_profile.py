import csv
from io import BytesIO
from pathlib import Path
import shutil
import sqlite3
import tempfile
import unittest
import uuid

from PIL import Image
from core import Store, export_csv
from company_profile import normalize_logo
from report_exports import export_excel, export_pdf, export_print_html


class CompanyProfileTests(unittest.TestCase):
    def setUp(self):
        self.folder=Path(tempfile.gettempdir())/('oasis-company-test-'+uuid.uuid4().hex)
        self.folder.mkdir()
        self.path=self.folder/'attendance.sqlite'
        self.store=Store(self.path)
        self.store.account('owner','password','admin',True)
        self.store.login('owner','password')
        image=Image.new('RGB',(120,60),'#D39340');buffer=BytesIO();image.save(buffer,format='JPEG')
        self.logo=buffer.getvalue()
        self.rows=[{'badge':'001','name':'Example','worked_minutes':480}]

    def tearDown(self):
        self.store.db.close();shutil.rmtree(self.folder)

    def test_defaults_and_database_isolation(self):
        self.assertEqual(self.store.company_profile(),{'company_name':'','branch':'','logo':b''})
        self.store.save_company_profile('شركة اختبار','الرياض',self.logo)
        other=Store(self.folder/'other.sqlite')
        try:self.assertEqual(other.company_profile()['company_name'],'')
        finally:other.db.close()

    def test_restart_backup_and_logo_removal(self):
        self.store.save_company_profile('Example Company','Jeddah',self.logo)
        self.store.db.close();self.store=Store(self.path)
        profile=self.store.company_profile()
        self.assertEqual(profile['company_name'],'Example Company')
        self.assertEqual(profile['branch'],'Jeddah')
        self.assertTrue(profile['logo'].startswith(b'\x89PNG'))
        backup=sqlite3.connect(self.folder/'backup.sqlite');self.store.db.backup(backup);backup.close()
        copied=Store(self.folder/'backup.sqlite')
        try:self.assertEqual(copied.company_profile(),profile)
        finally:copied.db.close()
        self.store.login('owner','password');self.store.save_company_profile('Example Company','Dammam',b'')
        self.assertEqual(self.store.company_profile()['logo'],b'')

    def test_admin_only(self):
        for role in ('manager','reports'):
            self.store.login('owner','password');self.store.account(role,'password',role)
            self.store.login(role,'password')
            with self.assertRaises(PermissionError):self.store.save_company_profile('Denied','',self.logo)
        self.assertEqual(self.store.company_profile()['company_name'],'')

    def test_invalid_image_and_names_do_not_change_saved_profile(self):
        self.store.save_company_profile('Before','Branch',self.logo)
        before=self.store.company_profile();audit=len(self.store.rows('SELECT * FROM audit'))
        for name,logo in [('After',b'not an image'),('x'*161,self.logo),('line\nbreak',self.logo)]:
            with self.assertRaises(ValueError):self.store.save_company_profile(name,'Branch',logo)
        self.assertEqual(self.store.company_profile(),before)
        self.assertEqual(len(self.store.rows('SELECT * FROM audit')),audit)

    def test_profile_save_rolls_back_if_audit_fails(self):
        self.store.db.execute("CREATE TRIGGER fail_profile_audit BEFORE INSERT ON audit BEGIN SELECT RAISE(ABORT,'test rollback'); END")
        self.store.db.commit()
        with self.assertRaises(sqlite3.IntegrityError):self.store.save_company_profile('Company','Branch',self.logo)
        self.assertEqual(self.store.company_profile()['company_name'],'')

    def test_logo_limits_and_proportions(self):
        with self.assertRaises(ValueError):normalize_logo(b'x'*(5*1024*1024+1))
        buffer=BytesIO();Image.new('RGB',(1200,300),'white').save(buffer,format='PNG')
        with Image.open(BytesIO(normalize_logo(buffer.getvalue()))) as logo:self.assertEqual(logo.size,(600,150))

    def test_reports_include_names_and_embedded_logo(self):
        from openpyxl import load_workbook
        profile={'company_name':'Example Company','branch':'Jeddah','logo':normalize_logo(self.logo)}
        for extension,exporter in [('xlsx',export_excel),('xls',export_excel),('pdf',export_pdf),('html',export_print_html)]:
            path=self.folder/('report.'+extension);exporter(path,self.rows,company=profile)
            self.assertGreater(path.stat().st_size,200)
        book=load_workbook(self.folder/'report.xlsx');sheet=book.active
        self.assertEqual(sheet['A3'].value,'Example Company');self.assertEqual(sheet['A4'].value,'Branch: Jeddah')
        self.assertEqual(sheet['A8'].value,'001');self.assertEqual(len(sheet._images),1);book.close()
        html=(self.folder/'report.html').read_text(encoding='utf-8')
        for text in ('Example Company','Jeddah','data:image/png;base64,'):self.assertIn(text,html)
        pdf=(self.folder/'report.pdf').read_bytes();self.assertIn(b'/Subtype /Image',pdf)
        xls=(self.folder/'report.xls').read_bytes()
        for text in (b'Example Company',b'Jeddah'):self.assertIn(text,xls)

    def test_company_text_is_not_html_or_excel_formula(self):
        from openpyxl import load_workbook
        profile={'company_name':'=1+1','branch':'<script>alert(1)</script>','logo':b''}
        export_excel(self.folder/'safe.xlsx',self.rows,company=profile)
        book=load_workbook(self.folder/'safe.xlsx');self.assertEqual(book.active['A3'].data_type,'s');book.close()
        export_print_html(self.folder/'safe.html',self.rows,company=profile)
        html=(self.folder/'safe.html').read_text(encoding='utf-8')
        self.assertNotIn('<script>',html);self.assertIn('&lt;script&gt;',html)
        export_csv(self.folder/'safe.csv',self.rows,company=profile)
        with open(self.folder/'safe.csv',encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        self.assertEqual(rows[0],['Company name',"'=1+1"])
        self.assertEqual(rows[1],['Branch',profile['branch']])
        self.assertEqual(rows[4][0],'001')


if __name__=='__main__':unittest.main()
