"""Report output adapters; export exactly the displayed report snapshot."""
from pathlib import Path
from datetime import datetime
from html import escape
import sys, os, base64

PRODUCER='This program produced by Oasis Digital Solutions'
COPYRIGHT='© 2026 Oasis Digital Solutions. All rights reserved.'

def asset(name):
    return Path(getattr(sys,'_MEIPASS',Path(__file__).parent))/'assets'/name

def check(rows):
    if not rows: raise ValueError('No rows to export.')

def export_excel(path,rows,title='Oasis Attend | Attendance report',translate=lambda x:x,arabic=False,summary=None):
    check(rows)
    columns=list(rows[0]); path=Path(path)
    if path.suffix.lower()=='.xls':
        import xlwt
        book=xlwt.Workbook(encoding='utf-8')
        heading=xlwt.easyxf('font: bold on, colour white; pattern: pattern solid, fore_colour dark_blue; align: wrap on, vert centre;')
        title_style=xlwt.easyxf('font: bold on, height 280;')
        text_style=xlwt.easyxf('align: vert top, wrap on;')
        for chunk,offset in enumerate(range(0,len(rows),65000),1):
            sheet=book.add_sheet('Report '+str(chunk));sheet.set_panes_frozen(True);sheet.set_horz_split_pos(4)
            sheet.write_merge(0,0,0,len(columns)-1,title,title_style)
            sheet.write_merge(1,1,0,len(columns)-1,'Oasis Digital Solutions | '+COPYRIGHT)
            for j,c in enumerate(columns):sheet.write(3,j,translate(c) if translate(c)!=c else c.replace('_',' ').title(),heading);sheet.col(j).width=256*(28 if c in ('name','first','last','stamp') else 20)
            sheet.row(3).height=650
            for i,r in enumerate(rows[offset:offset+65000],4):
                for j,c in enumerate(columns):
                    v=r.get(c,'');v=translate(v) if c=='status' else v
                    sheet.write(i,j,v if v is not None else '',text_style)
            if summary and offset==0:
                start=len(rows[offset:offset+65000])+5
                sheet.write(start,0,translate('Summary'),title_style)
                for i,(label,value) in enumerate(summary,start+1):
                    sheet.write(i,0,str(label),text_style);sheet.write(i,1,str(value),text_style)
            sheet.set_portrait(False);sheet.set_fit_width_to_pages(1);sheet.set_fit_height_to_pages(0);sheet.set_cols_right_to_left(arabic)
        book.save(str(path))
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image
        book=Workbook();sheet=book.active;sheet.title='Report';sheet.sheet_view.rightToLeft=arabic
        sheet.append([title]);sheet.append([COPYRIGHT]);sheet.append([]);sheet.append([translate(c) if translate(c)!=c else c.replace('_',' ').title() for c in columns]);sheet.freeze_panes='A5'
        for r in rows:
            sheet.append([translate(r.get(c,'')) if c=='status' else r.get(c,'') for c in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        for c in sheet[4]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='203E84');c.alignment=Alignment(wrap_text=True)
        sheet.row_dimensions[4].height=34;sheet['A1'].font=Font(size=17,bold=True)
        for j,c in enumerate(columns,1):sheet.column_dimensions[get_column_letter(j)].width=29 if c in ('name','first','last','stamp') else 21
        sheet.auto_filter.ref=f'A4:{get_column_letter(len(columns))}{sheet.max_row}';sheet.print_title_rows='1:4'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        if summary:
            sheet.append([])
            sheet.append([translate('Summary')])
            sheet[f'A{sheet.max_row}'].font=Font(bold=True)
            for label,value in summary:
                sheet.append([str(label),str(value)])
        sheet.oddFooter.center.text='Oasis Digital Solutions | Page &P of &N'
        book.save(path)
    else:raise ValueError('Choose .xls or .xlsx.')

def export_pdf(path,rows,title='Oasis Attend | Attendance report',translate=lambda x:x,arabic=False,summary=None):
    check(rows)
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,A3,landscape
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,Image
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font=Path(os.environ.get('WINDIR',r'C:\Windows'))/'Fonts/arial.ttf'
    if not font.exists():raise RuntimeError('Arial font is required for PDF export.')
    if 'ODS' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('ODS',str(font)))
    import arabic_reshaper
    from bidi.algorithm import get_display
    def shape(value):
        value=str(value if value is not None else '')
        import re
        value=re.sub(r'\d{4}-\d{2}-\d{2}',lambda m:'\u200e'+m.group()+'\u200e',value) if any('\u0600'<=c<='\u06ff' for c in value) else value
        return get_display(arabic_reshaper.reshape(value)) if any('\u0600'<=c<='\u06ff' for c in value) else value
    cols=list(rows[0]);size=landscape(A3 if len(cols)>10 else A4)
    doc=SimpleDocTemplate(str(path),pagesize=size,rightMargin=26,leftMargin=26,topMargin=28,bottomMargin=32,title=title,author='Oasis Digital Solutions')
    cell=ParagraphStyle('cell',fontName='ODS',fontSize=8.5,leading=12,alignment=2 if arabic else 0,wordWrap='CJK')
    head=ParagraphStyle('head',parent=cell,textColor=colors.white)
    def p(v,style=cell):return Paragraph(escape(shape(v)),style)
    weights=[2.0 if c in ('name','first','last','stamp') else 1.45 if c in ('department','shift','status') else 1.05 for c in cols]
    widths=[doc.width*w/sum(weights) for w in weights]
    table=Table([[p(translate(c) if translate(c)!=c else c.replace('_',' ').title(),head) for c in cols]]+[[p(translate(r.get(c,'')) if c=='status' else r.get(c,'')) for c in cols] for r in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#203E84')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F0F4FA')]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#203E84'))]))
    logo=asset('ods-logo.png');story=[]
    if logo.exists():story.append(Image(str(logo),width=155.25,height=51.75,hAlign='LEFT'))
    story += [Spacer(1,8),p(title,ParagraphStyle('title',parent=cell,fontSize=16,leading=22)),p('Oasis Digital Solutions')]
    if summary:
        # The same end-of-report figures shown on screen, kept with the report.
        head_style=ParagraphStyle('summary',parent=cell,fontSize=9,leading=12)
        pairs=[[p(str(label),head_style) for label,_ in summary],
               [p(str(value),ParagraphStyle('value',parent=cell,fontSize=11,leading=14)) for _,value in summary]]
        block=Table(pairs,colWidths=[doc.width/len(summary)]*len(summary),hAlign='LEFT')
        block.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#EDF1F8')),
                                   ('BOX',(0,0),(-1,-1),.6,colors.HexColor('#C3CEE4')),
                                   ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
                                   ('LEFTPADDING',(0,0),(-1,-1),7)]))
        story += [Spacer(1,10),block]
    story += [Spacer(1,12),table]
    def footer(c,d):
        c.setFont('ODS',8);c.setFillColor(colors.HexColor('#536079'));c.drawString(26,16,COPYRIGHT);c.drawRightString(size[0]-26,16,f'Page {d.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def export_print_html(path,rows,title='Oasis Attend | Attendance report',translate=lambda x:x,arabic=False,summary=None):
    check(rows);cols=list(rows[0]);logo=asset('ods-logo.png')
    image='<img alt="Oasis Digital Solutions" width="207" height="69" src="data:image/png;base64,'+base64.b64encode(logo.read_bytes()).decode()+'">' if logo.exists() else ''
    html='<!doctype html><html lang="'+('ar' if arabic else 'en')+'" dir="'+('rtl' if arabic else 'ltr')+'"><meta charset="utf-8"><title>'+escape(title)+'</title>'
    html+='''<style>body{font:12px Arial;margin:25px;color:#182b4d}h1{font-size:20px}table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{padding:7px;border:1px solid #ccd4df;overflow-wrap:anywhere;text-align:start}th{background:#203e84;color:white}tr:nth-child(even){background:#f0f4fa}thead{display:table-header-group}tr{break-inside:avoid}button{padding:10px 20px;margin:15px 0}.summary{display:flex;flex-wrap:wrap;gap:18px;background:#eef2f9;border:1px solid #c3cee4;padding:10px 14px;margin:10px 0}.summary div{display:flex;flex-direction:column}.summary b{font-size:15px}.summary span{font-size:11px;color:#425068}@page{size:landscape;margin:12mm}@media print{button{display:none}body{margin:0;font-size:9px}}</style>'''
    html+=image+'<h1>'+escape(title)+'</h1><p>'+escape(COPYRIGHT)+'</p>'
    if summary:
        html+='<div class="summary">'+''.join(
            '<div><b>'+escape(str(value))+'</b><span>'+escape(str(label))+'</span></div>'
            for label,value in summary)+'</div>'
    html+='<button onclick="window.print()">Print / طباعة</button><table><thead><tr>'+''.join('<th>'+escape(translate(c))+'</th>' for c in cols)+'</tr></thead><tbody>'
    for r in rows:html+='<tr>'+''.join('<td>'+escape(str(translate(r.get(c,'')) if c=='status' else r.get(c,'')))+'</td>' for c in cols)+'</tr>'
    Path(path).write_text(html+'</tbody></table></html>',encoding='utf-8')
